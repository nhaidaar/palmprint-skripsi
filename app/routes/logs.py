from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

router = APIRouter()


ALLOWED_LOG_STATUSES = {"ALLOWED", "DENIED"}
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_HEADERS = ("Time", "Name", "NIM", "Status", "Match %", "Duration", "Description")
EXCEL_COLUMN_WIDTHS = (20, 24, 16, 14, 12, 14, 36)


def _clean_status(status: str | None) -> str | None:
    if status in (None, ""):
        return None
    clean = status.upper()
    if clean not in ALLOWED_LOG_STATUSES:
        raise HTTPException(status_code=400, detail="status must be ALLOWED or DENIED")
    return clean


def _clean_date(value: str | None, name: str) -> str | None:
    if value in (None, ""):
        return None
    if len(value) != 10 or value[4] != "-" or value[7] != "-":
        raise HTTPException(status_code=400, detail=f"{name} must use YYYY-MM-DD")
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{name} must use YYYY-MM-DD") from exc


def _spreadsheet_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _safe_text(value, fallback: str = "—") -> str:
    if value in (None, ""):
        return fallback
    return _spreadsheet_safe(str(value))


def _filter_label(value: str | None) -> str:
    return _safe_text(value, "All")


def _excel_timestamp(value):
    if isinstance(value, datetime):
        return value
    if value in (None, ""):
        return "—"
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return _spreadsheet_safe(value)
    return value


def _build_logs_workbook(
    rows: list[dict],
    *,
    q: str | None,
    status: str | None,
    start_date: str | None,
    end_date: str | None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Access Logs"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    accent_fill = PatternFill(fill_type="solid", fgColor="595BD4")
    alternate_fill = PatternFill(fill_type="solid", fgColor="F8F9FC")
    allowed_fill = PatternFill(fill_type="solid", fgColor="E7F6EC")
    denied_fill = PatternFill(fill_type="solid", fgColor="FDEBEC")
    border_side = Side(style="thin", color="E5E7EB")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    sheet.merge_cells("A1:G1")
    title = sheet["A1"]
    title.value = "PalmGate Access Logs"
    title.fill = accent_fill
    title.font = Font(color="FFFFFF", bold=True, size=16)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells("A2:G2")
    summary = sheet["A2"]
    summary.value = (
        f"Search: {_filter_label(q)} | Status: {_filter_label(status)} | "
        f"From: {_filter_label(start_date)} | To: {_filter_label(end_date)}"
    )
    summary.font = Font(color="6B7280", italic=True, size=10)
    summary.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 22

    for column, header in enumerate(EXCEL_HEADERS, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.fill = accent_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
    sheet.row_dimensions[4].height = 24

    for column, width in enumerate(EXCEL_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    if not rows:
        sheet.merge_cells("A5:G5")
        empty = sheet["A5"]
        empty.value = "No matching logs"
        empty.font = Font(color="6B7280", italic=True)
        empty.alignment = Alignment(horizontal="center", vertical="center")
        empty.border = cell_border
        sheet.row_dimensions[5].height = 28
        sheet.auto_filter.ref = "A4:G4"
    else:
        for excel_row, row in enumerate(rows, start=5):
            values = (
                _excel_timestamp(row.get("timestamp")),
                _safe_text(row.get("matched_name"), "Unknown"),
                _safe_text(row.get("current_nim")),
                _safe_text(row.get("status")),
                row.get("similarity"),
                row.get("duration_ms"),
                _safe_text(row.get("description")),
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=excel_row, column=column, value=value)
                cell.border = cell_border
                cell.alignment = Alignment(vertical="center")
                if excel_row % 2 == 0:
                    cell.fill = alternate_fill

            sheet.cell(excel_row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
            sheet.cell(excel_row, 5).number_format = "0%"
            sheet.cell(excel_row, 6).number_format = '0 "ms"'

            status_cell = sheet.cell(excel_row, 4)
            if status_cell.value == "ALLOWED":
                status_cell.fill = allowed_fill
                status_cell.font = Font(color="237A3B", bold=True)
            elif status_cell.value == "DENIED":
                status_cell.fill = denied_fill
                status_cell.font = Font(color="B4232D", bold=True)

        sheet.auto_filter.ref = f"A4:G{sheet.max_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get("/api/logs/count")
async def get_logs_count(
    q: str | None = None,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
):
    from app.main import db
    return {
        "count": db.count_access_logs(
            q=q,
            status=_clean_status(status),
            start_date=_clean_date(start_date, "start_date"),
            end_date=_clean_date(end_date, "end_date"),
        )
    }


@router.get("/api/logs")
async def get_logs(
    limit: int = 20,
    offset: int = 0,
    q: str | None = None,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
):
    from app.main import db
    return db.get_access_logs(
        limit=limit,
        offset=offset,
        q=q,
        status=_clean_status(status),
        start_date=_clean_date(start_date, "start_date"),
        end_date=_clean_date(end_date, "end_date"),
    )


@router.get("/api/logs/export.xlsx")
async def export_logs_excel(
    q: str | None = None,
    status: str | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
):
    from app.main import db

    clean_status = _clean_status(status)
    clean_start_date = _clean_date(start_date, "start_date")
    clean_end_date = _clean_date(end_date, "end_date")
    rows = db.get_access_logs(
        limit=None,
        offset=0,
        q=q,
        status=clean_status,
        start_date=clean_start_date,
        end_date=clean_end_date,
    )
    workbook = _build_logs_workbook(
        rows,
        q=q,
        status=clean_status,
        start_date=clean_start_date,
        end_date=clean_end_date,
    )
    return Response(
        workbook,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f"attachment; filename={datetime.now():%Y%m%d-%H%M%S-%f}-palmgate.xlsx"},
    )
