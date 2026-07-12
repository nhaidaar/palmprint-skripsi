import re
from datetime import datetime
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


class FakeLogsDB:
    def __init__(self):
        self.list_calls = []
        self.count_calls = []
        self.rows = [
            {
                "id": 1,
                "timestamp": "2026-07-05 10:00:00",
                "status": "ALLOWED",
                "matched_name": "Alice",
                "current_nim": "A001",
                "similarity": 0.95,
                "duration_ms": 12,
                "description": "front door",
                "user_id": 7,
            }
        ]

    def get_access_logs(self, limit=20, offset=0, *, q=None, status=None, start_date=None, end_date=None):
        self.list_calls.append(
            {"limit": limit, "offset": offset, "q": q, "status": status, "start_date": start_date, "end_date": end_date}
        )
        return self.rows

    def count_access_logs(self, *, q=None, status=None, start_date=None, end_date=None):
        self.count_calls.append({"q": q, "status": status, "start_date": start_date, "end_date": end_date})
        return 1


def test_logs_route_forwards_filters(monkeypatch):
    import app.main as main

    fake_db = FakeLogsDB()
    monkeypatch.setattr(main, "db", fake_db)
    client = TestClient(app)

    response = client.get("/api/logs?limit=5&offset=10&q=alice&status=ALLOWED&start_date=2026-07-01&end_date=2026-07-05")

    assert response.status_code == 200
    assert response.json()[0]["current_nim"] == "A001"
    assert fake_db.list_calls == [
        {"limit": 5, "offset": 10, "q": "alice", "status": "ALLOWED", "start_date": "2026-07-01", "end_date": "2026-07-05"}
    ]


def test_logs_count_route_forwards_filters(monkeypatch):
    import app.main as main

    fake_db = FakeLogsDB()
    monkeypatch.setattr(main, "db", fake_db)
    client = TestClient(app)

    response = client.get("/api/logs/count?q=alice&status=DENIED&start_date=2026-07-01&end_date=2026-07-05")

    assert response.status_code == 200
    assert response.json() == {"count": 1}
    assert fake_db.count_calls == [
        {"q": "alice", "status": "DENIED", "start_date": "2026-07-01", "end_date": "2026-07-05"}
    ]


def test_logs_reject_invalid_status(monkeypatch):
    import app.main as main

    monkeypatch.setattr(main, "db", FakeLogsDB())
    client = TestClient(app)

    response = client.get("/api/logs?status=MAYBE")

    assert response.status_code == 400
    assert "status must be ALLOWED or DENIED" in response.json()["detail"]


def test_logs_reject_invalid_date_filters(monkeypatch):
    import app.main as main

    monkeypatch.setattr(main, "db", FakeLogsDB())
    client = TestClient(app)

    response = client.get("/api/logs?start_date=not-a-date")

    assert response.status_code == 400
    assert "start_date must use YYYY-MM-DD" in response.json()["detail"]


def test_logs_export_excel(monkeypatch):
    import app.main as main

    fake_db = FakeLogsDB()
    monkeypatch.setattr(main, "db", fake_db)
    client = TestClient(app)

    response = client.get("/api/logs/export.xlsx?q=alice&status=ALLOWED")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert re.fullmatch(
        r"attachment; filename=\d{8}-\d{6}-\d{6}-palmgate\.xlsx",
        response.headers["content-disposition"],
    )
    assert fake_db.list_calls[0] == {
        "limit": None,
        "offset": 0,
        "q": "alice",
        "status": "ALLOWED",
        "start_date": None,
        "end_date": None,
    }

    sheet = load_workbook(BytesIO(response.content))["Access Logs"]
    assert sheet["A1"].value == "PalmGate Access Logs"
    assert sheet["A2"].value == "Search: alice | Status: ALLOWED | From: All | To: All"
    assert [sheet.cell(4, column).value for column in range(1, 8)] == [
        "Time",
        "Name",
        "NIM",
        "Status",
        "Match %",
        "Duration",
        "Description",
    ]
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref == "A4:G5"
    assert sheet["A5"].value == datetime(2026, 7, 5, 10, 0, 0)
    assert sheet["B5"].value == "Alice"
    assert sheet["C5"].value == "A001"
    assert sheet["D5"].value == "ALLOWED"
    assert sheet["E5"].value == 0.95
    assert sheet["E5"].number_format == "0%"
    assert sheet["F5"].value == 12
    assert sheet["F5"].number_format == '0 "ms"'
    assert sheet["G5"].value == "front door"
    assert sheet["A4"].fill.fgColor.rgb[-6:] == "595BD4"
    assert sheet["D5"].fill.fgColor.rgb[-6:] == "E7F6EC"
    assert sheet.column_dimensions["G"].width == 36


def test_logs_export_excel_neutralizes_spreadsheet_formulas(monkeypatch):
    import app.main as main

    fake_db = FakeLogsDB()
    fake_db.rows = [{**fake_db.rows[0], "matched_name": "=cmd", "description": "+SUM(1,1)"}]
    monkeypatch.setattr(main, "db", fake_db)
    client = TestClient(app)

    response = client.get("/api/logs/export.xlsx")

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content))["Access Logs"]
    assert sheet["B5"].value == "'=cmd"
    assert sheet["G5"].value == "'+SUM(1,1)"
    assert sheet["B5"].data_type == "s"
    assert sheet["G5"].data_type == "s"


def test_logs_export_excel_handles_empty_results(monkeypatch):
    import app.main as main

    fake_db = FakeLogsDB()
    fake_db.rows = []
    monkeypatch.setattr(main, "db", fake_db)
    client = TestClient(app)

    response = client.get("/api/logs/export.xlsx")

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content))["Access Logs"]
    assert sheet["A5"].value == "No matching logs"
    assert "A5:G5" in {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert sheet.auto_filter.ref == "A4:G4"
